"""The web API, offline: projects are listed with their about.md, a run is serialized, an upstream failure is a
502 with a message, a dataset run streams every row and is saved, and saved runs report the same numbers as
core/report. run_experiment is faked; no keys, no spend."""

import importlib
import json
from collections import defaultdict, deque
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from core import limits, report
from core.projects import discover
from core.run import Result, Run

api = importlib.import_module("api.main")


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setattr(api, "RUNS", tmp_path)
    return TestClient(api.app)


def fake_result(label="billing"):
    return Result.of(Run("jev", "typesafe/jev", label, 0.93, 100.0, 50, 0, 0.000002),
                     Run("baseline", "fake-llm", "billing", None, 400.0, 60, 3, 0.00004),
                     Run("structured", "fake-llm", None, None, 450.0, 70, 6, 0.00005))


def test_every_project_is_listed_with_its_page(client):
    items = client.get("/api/projects").json()
    assert [p["id"] for p in items] == list(discover()) and len(items) == 24
    for p in items:
        d = client.get(f"/api/projects/{p['id']}").json()
        assert d["description"] and len(d["diagrams"]) == 2 and d["examples"]
        assert {x["heading"] for x in d["diagrams"]} == {"With Jev", "Without Jev"}
    assert client.get("/api/projects/99").status_code == 404


def test_a_run_is_serialized_with_titles_and_outcomes(client, monkeypatch):
    monkeypatch.setattr(discover()["01"], "run_experiment", lambda text: fake_result())
    body = client.post("/api/projects/01/run", json={"input": "I was charged twice"}).json()
    jev, _, structured = body["runs"]
    assert jev["label"] == "billing" and jev["title"] == "Jev-enhanced" and jev["confidence"] == 0.93
    assert structured["outcome"] == "(invalid)" and body["trace_url"] is None


def test_paid_runs_are_capped_per_visitor_and_for_everyone(client, monkeypatch):
    monkeypatch.setattr(discover()["01"], "run_experiment", lambda text: fake_result())
    monkeypatch.setattr(api.settings, "RUN_LIMIT_PER_IP_HOUR", 2)
    monkeypatch.setattr(api.settings, "RUN_LIMIT_PER_DAY", 3)
    monkeypatch.setattr(limits, "_hits", defaultdict(deque))
    post = lambda ip: client.post("/api/projects/01/run", json={"input": "x"}, headers={"x-forwarded-for": ip})
    assert [post("1.1.1.1").status_code for _ in range(3)] == [200, 200, 429]
    assert post("1.1.1.1").json()["detail"].startswith("Limit reached: 2 runs per hour")
    assert post("2.2.2.2").status_code == 200   # another visitor has their own hourly cap ...
    r = post("3.3.3.3")                          # ... but the day's 3 are now used, for everyone
    assert r.status_code == 429 and "per day, for everyone" in r.json()["detail"]


def test_an_upstream_failure_is_a_502_with_a_message(client, monkeypatch):
    def boom(text):
        raise RuntimeError("Server error '520 <none>' for url 'https://openrouter.ai/api/alpha/decisions'\nTraceback")
    monkeypatch.setattr(discover()["01"], "run_experiment", boom)
    r = client.post("/api/projects/01/run", json={"input": "x"})
    assert r.status_code == 502 and r.json()["error"].startswith("The run failed (RuntimeError): Server error '520")


def test_a_dataset_run_streams_every_row_saves_it_and_reports_the_core_numbers(client, monkeypatch, tmp_path):
    m = discover()["01"]
    calls = {"n": 0}

    def fake(text):
        calls["n"] += 1
        if calls["n"] == 3:
            raise RuntimeError("429 Too Many Requests")
        return fake_result(label="billing" if calls["n"] % 2 else "technical")
    monkeypatch.setattr(m, "run_experiment", fake)
    res = client.post("/api/projects/01/dataset")
    assert "no-transform" in res.headers["cache-control"]   # or Next's gzip holds every event until the end
    events = [e.split("\n", 1)[0].removeprefix("event: ") for e in res.text.strip().split("\n\n")]
    assert events == ["start", *["working", "row"] * len(m.DATASET), "done"]
    (saved,) = list((tmp_path / "01").glob("*.json"))
    rows = json.loads(saved.read_text(encoding="utf-8"))
    assert len(rows) == len(m.DATASET) and rows[2]["error"].startswith("RuntimeError: 429")

    listing = client.get("/api/projects/01/runs").json()
    assert listing[0]["file"] == saved.name and listing[0]["errors"] == 1
    rep = client.get(f"/api/projects/01/runs/{saved.name}").json()["report"]
    core_rows = [report.row_from_saved(r) for r in rows]
    for v in ("jev", "baseline", "structured"):
        want = report.summarize(core_rows, v)
        got = rep["summary"][v]
        assert (got["correct"], got["errors"], got["invalid"]) == (want["correct"], want["errors"], want["invalid"])
        assert got["labels"] == dict(want["labels"])
    assert rep["label_keys"][-1] == "(invalid)" and rep["wrong"]
    row = client.get(f"/api/projects/01/runs/{saved.name}/rows/{rep['wrong'][0]['idx']}").json()
    assert row["expected"] == m.DATASET[row["idx"]]["label"] and len(row["runs"]) == 3 and row["runs"][0]["title"] == "Jev-enhanced"
    assert client.get(f"/api/projects/01/runs/{saved.name}/rows/999").status_code == 404
    assert client.get("/api/projects/01/runs/..%2F..%2Fsecrets.json").status_code in (400, 404)


def test_saved_runs_diff_row_by_row(client, tmp_path):
    item = discover()["02"].DATASET[0]
    a = [report.row_to_saved(0, item, Result.of(Run("jev", "m", "urgent", None, 1, 1, 1, 0.0)), None)]
    b = [report.row_to_saved(0, item, Result.of(Run("jev", "m", "not_urgent", None, 1, 1, 1, 0.0)), None)]
    (tmp_path / "02").mkdir()
    for name, rows in (("2026-01-01_a.json", a), ("2026-01-02_b.json", b)):
        (tmp_path / "02" / name).write_text(json.dumps(rows), encoding="utf-8")
    d = client.get("/api/projects/02/diff", params={"a": "2026-01-01_a.json", "b": "2026-01-02_b.json"}).json()
    assert d["changed"] == [{"idx": 0, "variant": "jev", "text": item["text"], "expected": item["label"],
                             "before": "urgent", "after": "not_urgent", "q_before": None, "q_after": None}]


def test_the_sweep_is_in_the_report_for_projects_that_declare_one():
    m = discover()["02"]
    rows = [{"text": it["text"], "expected": it["label"], "error": None,
             "result": Result.of(Run("jev", "m", it["label"], None, 1, 1, 1, 0.0,
                                     probability=0.9 if it["label"] == "urgent" else 0.1))} for it in m.DATASET]
    rep = report.report(rows, m.LABELS, None, m.SWEEP)
    assert rep["sweep"]["best"]["fp"] == 0 and rep["sweep"]["best"]["fn"] == 0 and len(rep["sweep"]["points"]) == 21


def test_downloads_give_the_dataset_examples_and_prompts_but_never_config(client):
    import csv, io, zipfile
    m = discover()["01"]
    res = client.get("/api/projects/01/download/dataset.csv")
    assert res.status_code == 200 and "attachment" in res.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(res.text)))
    assert rows[0][:3] == ["idx", "text", "label"] and len(rows) == len(m.DATASET) + 1
    data = client.get("/api/projects/01/download/data.json").json()
    assert data["examples"] == m.EXAMPLES and len(data["dataset"]) == len(m.DATASET)
    z = zipfile.ZipFile(io.BytesIO(client.get("/api/projects/01/download/bundle.zip").content))
    names = z.namelist()
    assert {"jev-lab-01.xlsx", "data.json", "dataset.csv", "examples.json", "prompts_and_code/experiment.py",
            "prompts_and_code/about.md", "prompts_and_code/prompts.json"} <= set(names)
    assert not any("tests" in n or "config" in n or ".env" in n for n in names)
    assert client.get("/api/projects/01/download/..%2F..%2Fconfig%2F.env").status_code == 404
    assert client.get("/api/projects/01/download/secrets.txt").status_code == 404


def test_the_workbook_has_the_dataset_examples_and_every_captured_system_prompt(client):
    import io
    from openpyxl import load_workbook
    m = discover()["01"]
    res = client.get("/api/projects/01/download/workbook.xlsx")
    assert res.status_code == 200 and res.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames == ["Dataset", "Examples", "System prompts"]
    assert wb["Dataset"].max_row == len(m.DATASET) + 1 and wb["Examples"].max_row == len(m.EXAMPLES) + 1
    rows = client.get("/api/projects/01/prompts").json()["rows"]
    prompts = [r[4] for r in wb["System prompts"].iter_rows(min_row=2, values_only=True) if isinstance(r[0], int)]
    assert prompts == [r["prompt"] for r in rows] and any("billing" in p for p in prompts)


def test_every_project_has_captured_prompts_and_no_key_in_them():
    import re
    for nn, m in discover().items():
        p = Path(m.__file__).parent / "prompts.json"
        assert p.exists(), f"{nn}: run scripts/capture_prompts.py {nn}"
        text = p.read_text(encoding="utf-8")
        assert not re.search(r"sk-[A-Za-z0-9_-]{16,}|Bearer\s+\S{16,}", text), nn
        assert report.prompt_rows(json.loads(text)), nn


def test_the_capture_records_the_request_body_never_the_headers():
    import importlib.util
    import httpx
    spec = importlib.util.spec_from_file_location("capture_prompts", Path(__file__).parent.parent / "scripts" / "capture_prompts.py")
    cap = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cap)
    body = {"model": "m", "messages": [{"role": "system", "content": "SYS"}, {"role": "user", "content": "hi"}]}
    req = httpx.Request("POST", "https://openrouter.ai/api/v1/chat/completions", json=body,
                        headers={"Authorization": "Bearer sk-or-v1-secretsecretsecretsecret"})
    call = cap.record(req)
    assert call["system"] == ["SYS"] and call["user"] == ["hi"] and "secret" not in json.dumps(call)
    assert cap.record(httpx.Request("GET", "https://example.com/")) is None


def test_identical_calls_merge_into_one_prompt_row_with_a_count():
    q = {"kind": "jev", "model": "j", "questions": {"keep": {"type": "noul", "instructions": "relevant?"}}, "state": {}}
    rows = report.prompt_rows({"calls": [q, q, {"kind": "llm", "model": "f", "system": ["S"], "user": ["u"], "tools": [], "response_format": None}]})
    assert [(r["kind"], r["count"]) for r in rows] == [("Jev question", 2), ("LLM system prompt", 1)]


def test_presence_counts_live_tabs_and_unique_visitors(client, monkeypatch):
    presence = importlib.import_module("api.presence")
    monkeypatch.setattr(presence, "_live", {})
    monkeypatch.setattr(presence, "_seen", set())
    beat = lambda vid: client.post("/api/presence", json={"id": vid}).json()
    assert beat("visitor-aaaa") == {"live": 1, "visitors": 1}
    assert beat("visitor-aaaa") == {"live": 1, "visitors": 1}   # a second tab or beat is still one person
    assert beat("visitor-bbbb") == {"live": 2, "visitors": 2}
    presence._live["visitor-aaaa"] -= presence.LIVE_S + 1        # a went quiet: no longer live, still visited
    assert beat("visitor-bbbb") == {"live": 1, "visitors": 2}
    assert client.post("/api/presence", json={"id": "x"}).status_code == 422
