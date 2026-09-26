"""Jev Lab's web API: the same project contract the Streamlit app uses (run_experiment, EXAMPLES, DATASET, ...),
served as JSON for the Next.js UI in web/. Every number comes from core/report.py.

Run: uv run uvicorn api.main:app --port 8000"""

import csv
import dataclasses
import io
import json
import re
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from fastapi import Depends, FastAPI, HTTPException, Request
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response, StreamingResponse
from pydantic import BaseModel

from config import settings
from core import report
from core.limits import over_limit
from core.projects import ROOT, about_path, discover, traced
from core.run import Result

RUNS = ROOT / "runs"   # saved dataset runs, one folder per project (committed: the deployed History shows them)
_FILE = re.compile(r"^[\w.-]+\.json$")

app = FastAPI(title="Jev Lab API")
if settings.WEB_ORIGINS or settings.WEB_ORIGIN_REGEX:   # deployed: the paid POSTs go browser → API directly (see _limit)
    app.add_middleware(CORSMiddleware, allow_origins=settings.WEB_ORIGINS, allow_origin_regex=settings.WEB_ORIGIN_REGEX,
                       allow_methods=["GET", "POST"],
                       allow_headers=["content-type"])

def _limit(kind: str):
    """The spam caps (core/limits.py) on a paid POST. The browser calls the API directly for these, so
    X-Forwarded-For is the visitor, not Vercel's proxy."""
    def check(request: Request):
        ip = (request.headers.get("x-forwarded-for") or "").split(",")[0].strip() or getattr(request.client, "host", "?")
        if msg := over_limit(kind, ip):
            raise HTTPException(429, msg)
    return check


def _json(obj) -> JSONResponse:
    """Run.raw can hold anything a project stored; default=str keeps one odd value from failing a whole page."""
    return JSONResponse(json.loads(json.dumps(obj, default=str)))


def _project(nn: str):
    projects = discover()
    if nn not in projects:
        raise HTTPException(404, f"no project {nn}")
    return projects[nn]


def _result(result: Result) -> dict:
    return {"runs": [dataclasses.asdict(r) | {"title": report.title(r.variant), "outcome": report.outcome(r)}
                     for r in result.runs.values()], "trace_url": result.trace_url}


def _kind(runs: list[dict]) -> str:
    """How a variant decides: Jev, a model baseline, or free code (every run $0 by a "(no model)"-style model)."""
    if runs and runs[0]["variant"] == "jev":
        return "jev"
    if all((r["cost_usd"] or 0) == 0 and str(r["model"]).startswith("(") for r in runs):
        return "free"
    return "llm"


def _saved_files(nn: str) -> list[Path]:
    d = RUNS / nn
    return sorted(d.glob("*.json"), reverse=True) if d.exists() else []   # names start with a date: newest first


def _load(nn: str, name: str) -> list[dict]:
    if not _FILE.match(name):
        raise HTTPException(400, "bad run name")
    path = RUNS / nn / name
    if not path.exists():
        raise HTTPException(404, f"no run {name}")
    return json.loads(path.read_text(encoding="utf-8"))


def _headline(saved: list[dict]) -> dict:
    """Per variant: accuracy, pass rate, cost; what the overview plots."""
    rows = [report.row_from_saved(r) for r in saved]
    out = {}
    for v in report.variants(rows):
        s = report.summarize(rows, v)
        runs = [r["runs"][v] for r in saved if r["runs"] and v in r["runs"]]
        out[v] = {"title": report.title(v), "kind": _kind(runs), "accuracy": s["accuracy"], "correct": s["correct"],
                  "done": s["rows"] - s["errors"], "pass_rate": s["pass_rate"], "passed": s["passed"],
                  "graded": s["graded"], "total_cost": s["total_cost"], "cost_per_pass": s["cost_per_pass"],
                  "p50_ms": s["p50_ms"]}
    return out


def _run_meta(path: Path) -> dict:
    saved = json.loads(path.read_text(encoding="utf-8"))
    spend = sum((run["cost_usd"] or 0) + (run["raw"].get("grade_cost") or 0)
                for r in saved for run in (r["runs"] or {}).values())
    return {"file": path.name, "label": path.stem.split("_", 2)[-1].replace("-", " "),
            "date": path.stem[:10], "rows": len(saved), "errors": sum(bool(r.get("error")) for r in saved),
            "spend": spend, "headline": _headline(saved)}


def _summary(m, nn: str) -> dict:
    latest = next(iter(_saved_files(nn)), None)
    return {"id": nn, "title": m.TITLE, "name": m.TITLE.split("·", 1)[-1].strip(), "primitive": m.PRIMITIVE,
            "labels": m.LABELS, "examples": list(m.EXAMPLES), "rows": len(m.DATASET), "traced": traced(m),
            "sweep": getattr(m, "SWEEP", None), "safety": getattr(m, "SAFETY", None),
            "latest": _run_meta(latest) if latest else None}


@app.get("/api/projects")
def projects():
    return _json([_summary(m, nn) for nn, m in discover().items()])


@app.get("/api/projects/{nn}")
def project(nn: str):
    m = _project(nn)
    path = about_path(m)
    description, diagrams = report.parse_about(path.read_text(encoding="utf-8")) if path.exists() else ("", [])
    return _json(_summary(m, nn) | {"description": description, "examples": m.EXAMPLES,
                                    "diagrams": [{"heading": h, "mermaid": body} for h, body in diagrams]})


class RunIn(BaseModel):
    input: str | dict


@app.post("/api/projects/{nn}/run", dependencies=[Depends(_limit("run"))])
def run(nn: str, body: RunIn):
    m = _project(nn)
    try:
        result = m.run_experiment(body.input)
    except Exception as e:   # an upstream 5xx or a dropped connection: a message, not a stack trace
        return JSONResponse({"error": f"The run failed ({type(e).__name__}): {str(e).splitlines()[0][:200]}. "
                                      "Nothing was saved; run it again."}, status_code=502)
    return _json(_result(result))


def _sse(event: str, data) -> str:
    return f"event: {event}\ndata: {json.dumps(data, default=str)}\n\n"


@app.post("/api/projects/{nn}/dataset", dependencies=[Depends(_limit("dataset"))])
def dataset(nn: str):
    """Every row, one at a time (latency measures the model, not a traffic jam), streamed as it finishes.
    The file is rewritten after each row, so a closed tab never throws away rows already paid for."""
    m = _project(nn)
    RUNS.joinpath(nn).mkdir(parents=True, exist_ok=True)
    path = RUNS / nn / f"{datetime.now(timezone.utc):%Y-%m-%d_%H%M%S}_ui-run.json"

    def stream():
        saved = []
        yield _sse("start", {"rows": len(m.DATASET), "file": path.name})
        for i, item in enumerate(m.DATASET):
            yield _sse("working", {"idx": i, "text": item["text"]})   # what is running now: a row can take seconds
            try:
                result, error = m.run_experiment(item.get("input", item["text"])), None
            except Exception as e:   # one failed row must not throw away the rows already paid for
                result, error = None, f"{type(e).__name__}: {e}"[:300]
            row = report.row_to_saved(i, item, result, error)
            saved.append(row)
            path.write_text(json.dumps(saved, default=str, indent=1, ensure_ascii=False), encoding="utf-8")
            yield _sse("row", {"idx": i, "text": row["text"], "expected": row["expected"], "error": error,
                               "got": {v: report.outcome(r) for v, r in result.runs.items()} if result else None})
        yield _sse("done", {"file": path.name})
    # no-transform: Next's proxy gzips responses, and gzip buffers the whole stream until it ends
    return StreamingResponse(stream(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache, no-transform", "X-Accel-Buffering": "no"})


def _dataset_csv(m) -> str:
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["idx", "text", "label", "note", "input"])
    for i, item in enumerate(m.DATASET):
        inp = item.get("input")
        w.writerow([i, item["text"], item["label"], item.get("note") or "",
                    "" if inp is None else inp if isinstance(inp, str) else json.dumps(inp, ensure_ascii=False)])
    return out.getvalue()


def _data_json(m, nn: str) -> str:
    return json.dumps({"id": nn, "title": m.TITLE, "primitive": m.PRIMITIVE, "labels": m.LABELS,
                       "examples": m.EXAMPLES, "dataset": m.DATASET}, indent=1, ensure_ascii=False, default=str)


def _prompts(m) -> dict:
    path = Path(m.__file__).parent / "prompts.json"   # written by scripts/capture_prompts.py
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {"calls": [], "missing": True}


def _workbook(m, nn: str) -> bytes:
    """One Excel file: the dataset, the examples, and every system prompt the experiment sends (as captured)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    sheets = [("Dataset", ["idx", "text", "label", "note", "input"], [[i, it["text"], it["label"], it.get("note") or "",
               "" if it.get("input") is None else it["input"] if isinstance(it["input"], str) else json.dumps(it["input"], ensure_ascii=False)]
               for i, it in enumerate(m.DATASET)], [6, 70, 16, 40, 50]),
              ("Examples", ["name", "input"], [[k, v if isinstance(v, str) else json.dumps(v, ensure_ascii=False, indent=1)]
               for k, v in m.EXAMPLES.items()], [28, 90])]
    p = _prompts(m)
    rows = report.prompt_rows(p)
    sheets.append(("System prompts", ["#", "kind", "model", "times per run", "prompt (verbatim)", "user message (captured example)"],
                   [[r["n"], r["kind"], r["model"], r["count"], r["prompt"], r["user"]] for r in rows], [5, 18, 28, 12, 100, 60]))
    for i, (title, head, body, widths) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(head)
        for row in body:
            ws.append(row)
        for c in ws[1]:
            c.font = Font(bold=True)
        for col, w in zip("ABCDEF", widths):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
    if rows:
        ws.append([])
        ws.append([f"Captured from a real run of example '{p.get('example')}' on {p.get('captured')} "
                   "(scripts/capture_prompts.py): exactly what was sent to each model."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@app.get("/api/projects/{nn}/prompts")
def prompts(nn: str):
    p = _prompts(_project(nn))
    return _json({"example": p.get("example"), "captured": p.get("captured"), "missing": p.get("missing", False),
                  "rows": report.prompt_rows(p)})


def _bundle(m, nn: str) -> bytes:
    """Data plus the prompts: the Jev questions, baseline prompts and policies live in the project's own .py files.
    Only files inside the project folder, never tests/ or config/ (keys live in config/.env)."""
    folder = Path(m.__file__).parent
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"jev-lab-{nn}.xlsx", _workbook(m, nn))   # dataset + examples + system prompts, one file
        z.writestr("data.json", _data_json(m, nn))
        z.writestr("dataset.csv", _dataset_csv(m))
        z.writestr("examples.json", json.dumps(m.EXAMPLES, indent=1, ensure_ascii=False, default=str))
        for f in sorted(folder.glob("*")):
            if f.is_file() and (f.suffix in (".py", ".md") or f.name == "prompts.json"):
                z.write(f, f"prompts_and_code/{f.name}")
    return buf.getvalue()


DOWNLOADS = {"dataset.csv": "text/csv; charset=utf-8", "data.json": "application/json", "bundle.zip": "application/zip",
             "workbook.xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}


@app.get("/api/projects/{nn}/download/{kind}")
def download(nn: str, kind: str):
    m = _project(nn)
    if kind not in DOWNLOADS:
        raise HTTPException(404, f"no download {kind}")
    body = {"dataset.csv": lambda: _dataset_csv(m), "data.json": lambda: _data_json(m, nn), "bundle.zip": lambda: _bundle(m, nn),
            "workbook.xlsx": lambda: _workbook(m, nn)}[kind]()
    name = f"jev-lab-{nn}-{Path(m.__file__).parent.name[3:]}-{kind}".replace("-bundle.zip", ".zip").replace("-workbook.xlsx", ".xlsx")
    return Response(body, media_type=DOWNLOADS[kind], headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/projects/{nn}/runs")
def runs(nn: str):
    _project(nn)
    return _json([_run_meta(p) for p in _saved_files(nn)])


@app.get("/api/projects/{nn}/runs/{name}")
def saved_run(nn: str, name: str, cost_fn: float = 5.0, cost_fp: float = 1.0):
    m = _project(nn)
    saved = _load(nn, name)
    rows = [report.row_from_saved(r) for r in saved]
    return _json({"meta": _run_meta(RUNS / nn / name),
                  "report": report.report(rows, m.LABELS, getattr(m, "SAFETY", None), getattr(m, "SWEEP", None),
                                          cost_fn, cost_fp)})


@app.get("/api/projects/{nn}/runs/{name}/rows/{idx}")
def saved_row(nn: str, name: str, idx: int):
    """One saved row with every variant's full run, shaped like a live run: inspect any past answer for free."""
    _project(nn)
    row = next((r for r in _load(nn, name) if r["idx"] == idx), None)
    if row is None:
        raise HTTPException(404, f"no row {idx}")
    result = report.row_from_saved(row)["result"]
    return _json({"idx": idx, "text": row["text"], "expected": row["expected"], "input": row.get("input"),
                  "error": row.get("error"), **(_result(result) if result else {"runs": [], "trace_url": None})})


@app.get("/api/projects/{nn}/diff")
def diff(nn: str, a: str, b: str):
    _project(nn)
    return _json(report.diff_runs(_load(nn, a), _load(nn, b)))
